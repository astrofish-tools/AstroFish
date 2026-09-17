from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


if len(sys.argv) != 3:
    raise SystemExit("usage: build_fixed_docx.py <rendered-pages-dir> <output.docx>")

pages_dir = Path(sys.argv[1])
output = Path(sys.argv[2])
pages = sorted(pages_dir.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
if not pages:
    raise SystemExit("no rendered pages found")

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(0.38)
section.left_margin = section.right_margin = Cm(0.3)
section.header_distance = section.footer_distance = Cm(0)

normal = doc.styles["Normal"]
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1

first = doc.add_paragraph()
for index, page in enumerate(pages):
    paragraph = first if index == 0 else doc.add_paragraph()
    if index > 0:
        paragraph.paragraph_format.page_break_before = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(page), width=Cm(19.95), height=Cm(28.21))

doc.core_properties.title = "AstroFish 1.0 使用说明 / User Guide"
doc.core_properties.author = "yxy"
doc.core_properties.subject = "Star View & Long Exposure Intervalometer for Sony α Cameras"
doc.save(output)
print(output)
