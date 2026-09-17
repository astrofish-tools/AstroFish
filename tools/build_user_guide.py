from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "deliverables" / "AstroFish-1.0.0"
MD = OUT / "AstroFish-1.0.0-User-Guide-ZH-EN.md"
PNG = OUT / "AstroFish-1.0.0-Interface.png"
DOCX = OUT / "AstroFish-1.0.0-User-Guide-ZH-EN.docx"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc" if bold else "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for item in candidates:
        if Path(item).exists():
            return ImageFont.truetype(item, size)
    return ImageFont.load_default()


def create_diagram():
    im = Image.new("RGB", (1800, 1020), "#080b11")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((30, 30, 1770, 990), radius=30, outline="#34445f", width=3, fill="#080b11")
    d.text((70, 60), "AstroFish 1.0", font=font(48, True), fill="#f2f5fa")
    d.text((1170, 74), "Sony α Camera App · by yxy", font=font(27), fill="#8aa0bf")
    box = (70, 145, 1730, 735)
    d.rounded_rectangle(box, radius=20, outline="#536783", width=3, fill="#0b1830")
    stars = [(130,205,4),(225,280,3),(340,190,5),(465,248,3),(620,190,4),(735,300,5),(890,218,3),(1010,285,4),(1160,180,4),(1305,250,3),(1490,205,5),(1620,300,3)]
    for x,y,r in stars: d.ellipse((x-r,y-r,x+r,y+r), fill="#dcecff")
    d.polygon([(70,625),(240,520),(370,590),(520,470),(700,600),(875,450),(1040,580),(1210,430),(1420,585),(1580,500),(1730,555),(1730,735),(70,735)], fill="#05080d")
    d.text((102, 175), "STAR VIEW", font=font(31, True), fill="#ffb14a")
    d.rounded_rectangle((95, 555, 1315, 710), radius=14, fill="#000000", outline="#25364d", width=2)
    d.text((125, 575), "[SS 15s]     F 2.8     ISO 1600", font=font(34), fill="#ffffff")
    d.text((125, 625), "SHOTS 20     DELAY 2s     INTERVAL 1s", font=font(28), fill="#d7dde8")
    d.text((125, 670), "C1 STAR/CLEAN · C2/AEL MAG · FN SELECT · PLAY REVIEW · MENU BULB", font=font(20), fill="#9bb2d0")
    d.rounded_rectangle((1360, 555, 1700, 710), radius=14, fill="#05070b", outline="#60728e", width=2)
    d.text((1380, 570), "LIVE HISTOGRAM", font=font(20), fill="#a9bad2")
    pts = [(1380+i*19, 685-v) for i,v in enumerate([4,8,14,25,50,70,45,25,18,35,66,94,73,45,24,15])]
    d.line(pts, fill="#edf3fb", width=4)
    d.line((1680,595,1680,690), fill="#ef4d4d", width=7)
    labels = [
        (80,790,"1", "实时取景 / Live view", "全屏构图与手动对焦"),
        (625,790,"2", "曝光与序列 / Controls", "快门、光圈、ISO、张数、延迟、间隔"),
        (1225,790,"3", "直方图 / Histogram", "实时辅助；回放确认最终照片"),
    ]
    for x,y,n,title,body in labels:
        d.ellipse((x,y,x+56,y+56), fill="#ff9b34")
        d.text((x+18,y+8), n, font=font(30, True), fill="#111722")
        d.text((x+75,y-2), title, font=font(27, True), fill="#f2f5fa")
        d.text((x+75,y+37), body, font=font(21), fill="#9eacc0")
    d.rounded_rectangle((80, 885, 1720, 960), radius=14, fill="#111722", outline="#26344a", width=2)
    d.text((105, 906), "C1：普通参数 → 普通纯净 → Star View 参数 → Star View 纯净 → 普通参数", font=font(26, True), fill="#ffb14a")
    im.save(PNG)


def set_run_font(run, size=None, bold=None, color=None, mono=False):
    name = "Menlo" if mono else "Arial"
    east = "Heiti TC"
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_inline(paragraph, text, size=9.4):
    token = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for match in token.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]), size=size)
        raw = match.group(0)
        if raw.startswith("**"):
            set_run_font(paragraph.add_run(raw[2:-2]), size=size, bold=True)
        else:
            run = paragraph.add_run(raw[1:-1])
            set_run_font(run, size=size-0.4, color=(196, 91, 39), mono=True)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.55)
    sec.left_margin, sec.right_margin = Cm(1.7), Cm(1.7)

    styles = doc.styles
    styles["Normal"].font.size = Pt(9.4)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for style_name, size, color in [("Title", 28, (19,43,75)), ("Heading 1", 18, (19,43,75)), ("Heading 2", 12.5, (201,83,34)), ("Heading 3", 10.5, (19,43,75))]:
        s = styles[style_name]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti TC")
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(*color)
        s.font.bold = True

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("AstroFish 1.0  ·  Sony α Camera App  ·  yxy"), size=8, color=(105,118,137))
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("AstroFish 1.0 User Guide  |  "), size=8, color=(105,118,137))
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); footer._p.append(fld)

    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    rule_count = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1; continue
        if line == "---":
            rule_count += 1
            doc.add_page_break()
            i += 1; continue
        if line.startswith("!["):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(PNG), width=Inches(6.55))
            i += 1; continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[-: |]+\|$", lines[i+1]):
            rows = []
            rows.append([c.strip() for c in line.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")]); i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for rr, row in enumerate(rows):
                for cc, value in enumerate(row):
                    cell = table.cell(rr, cc)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    add_inline(p, value, size=8.1)
                    if rr == 0:
                        shade(cell, "17365D")
                        for run in p.runs: run.font.color.rgb = RGBColor(255,255,255); run.bold = True
                    elif rr % 2 == 0: shade(cell, "EEF3F8")
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if line.startswith("# "):
            text = line[2:]
            p = doc.add_paragraph(style="Title" if i == 0 else "Heading 1")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(text)
            i += 1; continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
            i += 1; continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
            i += 1; continue
        m = re.match(r"^[-] (.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(1))
            i += 1; continue
        m = re.match(r"^(\d+)\. (.*)$", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(2)
            set_run_font(p.add_run(m.group(1) + ". "), size=9.4, bold=True, color=(201,83,34))
            add_inline(p, m.group(2))
            i += 1; continue
        p = doc.add_paragraph()
        if i < 6:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(7)
            add_inline(p, line.replace("  ", ""), size=10.5)
        else:
            add_inline(p, line.replace("  ", ""))
        i += 1

    doc.core_properties.title = "AstroFish 1.0 使用说明 / User Guide"
    doc.core_properties.author = "yxy"
    doc.core_properties.subject = "Star View & Long Exposure Intervalometer for Sony α Cameras"
    doc.save(DOCX)


if __name__ == "__main__":
    create_diagram()
    build_docx()
    print(DOCX)
